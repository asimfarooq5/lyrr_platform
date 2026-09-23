"""
Admin portal routes - full CRUD for books, users, subscriptions, categories, analytics
Serves Jinja2 templates with server-side rendering
"""

from fastapi import APIRouter, Depends, HTTPException, Request, Form, UploadFile, File
from fastapi.responses import RedirectResponse, HTMLResponse, Response
from fastapi.templating import Jinja2Templates
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, desc, or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import selectinload
from typing import Optional
from datetime import datetime, timedelta, timezone
import uuid
import os
import io
import logging
import shutil
import zipfile
import xml.etree.ElementTree as ET
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)

from app.core.config import settings
from app.core.database import get_db
from app.core.rate_limit import limiter
from app.core.security import verify_password, create_access_token, decode_token, get_password_hash
from app.core.csrf import generate_csrf_token, verify_csrf_token, revoke_csrf_tokens
from app.models.user import User
from app.models.book import Book, BookMedia, BookStatus, Chapter
from app.models.content import Category, Author, BookCategory, SubscriptionPlan, UserSubscription, Payment
from app.models.reading_session import ReadingSession

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", ".."))
STORAGE_DIR = os.path.join(BASE_DIR, "storage")
AUDIO_DIR = os.path.join(STORAGE_DIR, "audio")
COVERS_DIR = os.path.join(STORAGE_DIR, "covers")
os.makedirs(AUDIO_DIR, exist_ok=True)
os.makedirs(COVERS_DIR, exist_ok=True)

# ---- File validation constants ----
ALLOWED_AUDIO_EXTS = {".mp3", ".m4a", ".wav", ".flac", ".aac", ".ogg"}
ALLOWED_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp"}


def format_number(value):
    try:
        return "{:,.0f}".format(float(value))
    except (ValueError, TypeError):
        return str(value)


def _find_el(parent, prefixed: str, plain: str, ns: dict):
    """Find a child element by a namespaced name, falling back to the bare name.

    NOTE: must not use `a or b` on the results. An ElementTree Element with no
    child elements (e.g. `<dc:publisher>Foo</dc:publisher>`) is falsy, so
    `find(prefixed) or find(plain)` silently returns None even on a hit.
    """
    el = parent.find(prefixed, ns)
    if el is not None:
        return el
    return parent.find(plain)


def _make_soup(html: str):
    """Parse HTML, preferring lxml but degrading to the stdlib parser.

    lxml is faster and more forgiving, but it is a compiled dependency that
    may be missing on some platforms — an EPUB import must never hard-fail
    just because the optional parser is unavailable.
    """
    try:
        return BeautifulSoup(html, "lxml")
    except Exception:
        return BeautifulSoup(html, "html.parser")


def _parse_price(value: str) -> Optional[float]:
    """Parse a price form field. Blank -> None (free); invalid -> None (free).

    Never raises: a typo in the admin form must not 500 the request.
    """
    text = (value or "").strip()
    if not text:
        return None
    try:
        price = float(text)
    except (ValueError, TypeError):
        return None
    return price if price >= 0 else None

templates.env.filters["format_number"] = format_number


def cover_url(cover_path: Optional[str], updated_at=None) -> Optional[str]:
    """Append a cache-busting version to a cover URL.

    Cover files keep the same path when replaced, so without this the browser
    (and the Flutter image cache) would keep showing the old artwork after an
    admin uploads a new one.
    """
    if not cover_path:
        return cover_path
    version = int(updated_at.timestamp()) if updated_at else 1
    separator = "&" if "?" in cover_path else "?"
    return f"{cover_path}{separator}v={version}"


# Expose helpers to every template (upload size limits, cover cache-busting).
templates.env.globals["settings"] = settings
templates.env.globals["cover_url"] = cover_url


async def _get_admin(request: Request, db: AsyncSession) -> Optional[User]:
    token = request.cookies.get("admin_token")
    if not token:
        return None
    try:
        payload = decode_token(token)
        if payload and payload.get("type") == "admin":
            result = await db.execute(select(User).where(User.id == payload["sub"]))
            user = result.scalar_one_or_none()
            if user and user.is_admin:
                return user
    except Exception:
        pass
    return None


async def _inject_csrf(context: dict, request: Request) -> dict:
    """Inject CSRF token into template context"""
    if settings.CSRF_ENABLED:
        session_id = request.cookies.get("admin_token", "") or "anonymous"
        context["csrf_token"] = await generate_csrf_token(session_id)
    else:
        context["csrf_token"] = ""
    return context


async def _verify_csrf(request: Request) -> bool:
    """Verify CSRF token from form data against stored session token"""
    if not settings.CSRF_ENABLED:
        return True
    form = await request.form()
    token = form.get("csrf_token", "")
    session_id = request.cookies.get("admin_token", "")
    if not await verify_csrf_token(session_id, token):
        raise HTTPException(status_code=403, detail="CSRF validation failed. Please submit the form again.")
    return True


async def _validate_upload(
    file: UploadFile,
    allowed_exts: set,
    max_size_mb: int,
) -> tuple[str, bytes]:
    """Validate and sanitize an uploaded file. Returns (safe_filename, content)."""
    # Check extension
    original_name = file.filename or ""
    ext = os.path.splitext(original_name)[1].lower()
    if ext not in allowed_exts:
        raise HTTPException(
            status_code=400,
            detail=f"File type '{ext}' is not allowed. Allowed: {', '.join(sorted(allowed_exts))}",
        )

    # Read content with a size cap
    max_bytes = max_size_mb * 1024 * 1024
    content = await file.read(max_bytes + 1)

    # Check size
    if len(content) > max_bytes:
        raise HTTPException(
            status_code=400,
            detail=f"File too large ({len(content) / 1024 / 1024:.1f} MB). Maximum: {max_size_mb} MB",
        )

    # Sanitize filename — UUID + extension only
    safe_name = f"{uuid.uuid4().hex}{ext}"
    return safe_name, content


# ===== AUTH =====

@router.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    return templates.TemplateResponse(request, "admin/login.html", {"request": request, "error": None})

@router.post("/login")
@limiter.limit("10/minute")
async def login(request: Request, username: str = Form(...), password: str = Form(...),
                db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(User).where(User.email == username))
    user = result.scalar_one_or_none()
    if not user or not user.is_admin or not verify_password(password, user.hashed_password):
        return templates.TemplateResponse(request, "admin/login.html", {"request": request, "error": "Invalid credentials"})
    token = create_access_token(data={"sub": user.id, "type": "admin"})
    resp = RedirectResponse(url="/admin/dashboard", status_code=302)
    resp.set_cookie(key="admin_token", value=token, httponly=True, max_age=3600, path="/admin",
                    samesite="lax")
    return resp

@router.post("/logout")
async def logout(request: Request, _: bool = Depends(_verify_csrf)):
    """Log out of the admin portal"""
    session_id = request.cookies.get("admin_token", "")
    await revoke_csrf_tokens(session_id)
    resp = RedirectResponse(url="/admin/login", status_code=302)
    resp.delete_cookie(key="admin_token", path="/admin")
    return resp

@router.get("/dashboard", response_class=HTMLResponse)
async def dashboard(request: Request, db: AsyncSession = Depends(get_db)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    stats = {
        "users": (await db.scalar(select(func.count(User.id)))) or 0,
        "books": (await db.scalar(select(func.count(Book.id)))) or 0,
        "published": (await db.scalar(
            select(func.count(Book.id)).where(Book.status == BookStatus.PUBLISHED)
        )) or 0,
        "subscriptions": (await db.scalar(
            select(func.count(UserSubscription.id)).where(UserSubscription.status == "active")
        )) or 0,
        "payments_count": (await db.scalar(
            select(func.count(Payment.id)).where(Payment.status == "completed")
        )) or 0,
    }

    # Revenue totals (kept as SQL aggregates so they scale past a demo dataset).
    stats["revenue"] = int(await db.scalar(
        select(func.coalesce(func.sum(Payment.amount), 0))
        .where(Payment.status == "completed")
    ) or 0)
    since_30 = datetime.now(timezone.utc) - timedelta(days=30)
    stats["revenue_month"] = int(await db.scalar(
        select(func.coalesce(func.sum(Payment.amount), 0))
        .where(Payment.status == "completed", Payment.completed_at >= since_30)
    ) or 0)
    stats["sales_month"] = int(await db.scalar(
        select(func.count(Payment.id))
        .where(Payment.status == "completed", Payment.completed_at >= since_30)
    ) or 0)

    # 14-day revenue trend for the sparkline chart.
    trend_rows = (await db.execute(
        select(
            func.date(Payment.completed_at).label("day"),
            func.coalesce(func.sum(Payment.amount), 0).label("total"),
        )
        .where(Payment.status == "completed", Payment.completed_at >= since_30)
        .group_by(func.date(Payment.completed_at))
        .order_by(func.date(Payment.completed_at))
    )).all()
    by_day = {str(r.day): float(r.total) for r in trend_rows}
    trend = []
    for offset in range(13, -1, -1):
        day = (datetime.now(timezone.utc) - timedelta(days=offset)).date()
        trend.append({"day": day.strftime("%d %b"), "value": by_day.get(str(day), 0.0)})
    stats["trend"] = trend
    stats["trend_max"] = max((p["value"] for p in trend), default=0.0) or 1.0

    # Recent payments + best sellers for the activity panels.
    recent = (await db.execute(
        select(Payment, User.email)
        .join(User, User.id == Payment.user_id, isouter=True)
        .order_by(Payment.created_at.desc())
        .limit(6)
    )).all()
    recent_payments = [
        {
            "email": email or "—",
            "amount": p.amount,
            "currency": p.currency,
            "method": p.method,
            "status": p.status,
            "description": p.description,
            "created_at": p.created_at,
        }
        for p, email in recent
    ]

    top_books = (await db.execute(
        select(Book.id, Book.title, Book.author, Book.cover_url, Book.updated_at)
        .where(Book.status == BookStatus.PUBLISHED)
        .order_by(Book.sales_count.desc(), Book.created_at.desc())
        .limit(5)
    )).all()

    ctx = {
        "request": request,
        "page": "dashboard",
        "stats": stats,
        "recent_payments": recent_payments,
        "top_books": top_books,
    }
    return templates.TemplateResponse(request, "admin/dashboard.html", await _inject_csrf(ctx, request))

# ===== BOOKS =====

@router.get("/books", response_class=HTMLResponse)
async def books_page(request: Request, db: AsyncSession = Depends(get_db)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(Book).options(selectinload(Book.media)).order_by(Book.created_at.desc()).limit(100))
    books = result.scalars().all()
    ctx = {"request": request, "page": "books", "books": books}
    return templates.TemplateResponse(request, "admin/books.html", await _inject_csrf(ctx, request))

@router.post("/books")
async def create_book(request: Request, title: str = Form(...), author: str = Form(""),
                      description: str = Form(""), language: str = Form("en"),
                      price: str = Form(""),
                      db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    book = Book(id=str(uuid.uuid4()), title=title, author=author or "Unknown",
                description=description, language=language,
                price=_parse_price(price),
                status=BookStatus.DRAFT)
    db.add(book)
    await db.commit()
    return RedirectResponse(url=f"/admin/books/{book.id}", status_code=302)

@router.get("/books/{book_id}")
async def edit_book_page(book_id: str, request: Request, db: AsyncSession = Depends(get_db)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(Book).options(selectinload(Book.media)).where(Book.id == book_id))
    book = result.scalar_one_or_none()
    if not book:
        return RedirectResponse(url="/admin/books")
    categories = (await db.execute(select(Category).order_by(Category.name))).scalars().all()
    current_category = (await db.execute(
        select(BookCategory.category_id).where(BookCategory.book_id == book_id)
    )).scalar_one_or_none()
    chapter_count = await db.scalar(
        select(func.count(Chapter.id)).where(Chapter.book_id == book_id)
    ) or 0
    ctx = {"request": request, "page": "books", "book": book,
           "categories": categories, "current_category_id": current_category,
           "chapter_count": chapter_count}
    return templates.TemplateResponse(request, "admin/book_detail.html", await _inject_csrf(ctx, request))

@router.post("/books/{book_id}")
async def update_book(book_id: str, request: Request, title: str = Form(...),
                      author: str = Form(""), description: str = Form(""),
                      book_type: str = Form("fiction"),
                      price: str = Form(""),
                      category_id: str = Form(""),
                      status: str = Form("draft"),
                      db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(Book).where(Book.id == book_id))
    book = result.scalar_one_or_none()
    if not book:
        return RedirectResponse(url="/admin/books")
    book.title = title
    book.author = author or book.author
    book.description = description
    book.book_type = book_type
    book.price = _parse_price(price)
    if status in (BookStatus.DRAFT.value, BookStatus.PUBLISHED.value, BookStatus.ARCHIVED.value):
        book.status = BookStatus(status)
        if book.status == BookStatus.PUBLISHED and not book.published_at:
            book.published_at = datetime.utcnow()

    # Pricing & Publishing (KDP-style): one category per book, replace on save.
    await db.execute(
        BookCategory.__table__.delete().where(BookCategory.book_id == book_id)
    )
    if category_id.strip():
        db.add(BookCategory(book_id=book_id, category_id=category_id))

    await db.commit()
    return RedirectResponse(url=f"/admin/books/{book_id}", status_code=302)

@router.post("/books/{book_id}/cover")
async def upload_cover(book_id: str, request: Request, file: UploadFile = File(...),
                       db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    
    safe_name, content = await _validate_upload(file, ALLOWED_IMAGE_EXTS, settings.MAX_COVER_SIZE_MB)
    filename = f"{book_id}-{safe_name}"
    path = os.path.join(COVERS_DIR, filename)
    with open(path, "wb") as f:
        f.write(content)

    result = await db.execute(select(Book).where(Book.id == book_id))
    book = result.scalar_one_or_none()
    if not book:
        return RedirectResponse(url="/admin/books")
    book.cover_url = f"/media/covers/{filename}"
    await db.commit()
    return RedirectResponse(url=f"/admin/books/{book_id}", status_code=302)

@router.post("/books/{book_id}/audio")
async def upload_audio(book_id: str, request: Request, file: UploadFile = File(...),
                       db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    
    safe_name, content = await _validate_upload(file, ALLOWED_AUDIO_EXTS, settings.MAX_UPLOAD_SIZE_MB)
    filename = f"{book_id}-{safe_name}"
    path = os.path.join(AUDIO_DIR, filename)
    with open(path, "wb") as f:
        f.write(content)

    media = BookMedia(
        id=str(uuid.uuid4()), book_id=book_id,
        audio_url=f"/media/audio/{filename}",
        format=os.path.splitext(safe_name)[1].lstrip("."),
        size_bytes=len(content),
        is_encrypted=False,
    )
    db.add(media)
    await db.commit()
    return RedirectResponse(url=f"/admin/books/{book_id}", status_code=302)

@router.post("/books/{book_id}/audio/{media_id}/delete")
async def delete_audio(book_id: str, media_id: str, request: Request,
                        db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(BookMedia).where(BookMedia.id == media_id, BookMedia.book_id == book_id))
    media = result.scalar_one_or_none()
    if media:
        fname = os.path.basename(media.audio_url) if media.audio_url else ""
        fpath = os.path.join(AUDIO_DIR, fname)
        if os.path.exists(fpath):
            os.remove(fpath)
        await db.delete(media)
        await db.commit()
    return RedirectResponse(url=f"/admin/books/{book_id}", status_code=302)

@router.post("/books/{book_id}/delete")
async def delete_book_route(book_id: str, request: Request,
                             db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(Book).where(Book.id == book_id))
    book = result.scalar_one_or_none()
    if book:
        await db.delete(book)
        await db.commit()
    return RedirectResponse(url="/admin/books", status_code=302)


@router.post("/books/{book_id}/upload-epub")
async def upload_epub(book_id: str, request: Request, file: UploadFile = File(...),
                      db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    """Upload an EPUB file and extract chapters/content"""
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")

    MAX_EPUB_SIZE = 50 * 1024 * 1024  # 50 MB hard cap
    content = await file.read(MAX_EPUB_SIZE + 1)
    if len(content) > MAX_EPUB_SIZE:
        return RedirectResponse(url=f"/admin/books/{book_id}?error=epub_too_large", status_code=302)

    epub_dir = os.path.join(BASE_DIR, "storage", "epub_temp", uuid.uuid4().hex)
    os.makedirs(epub_dir, exist_ok=True)

    def _safe_join(base: str, *parts: str) -> str:
        """Join path parts and ensure the result stays inside base (zip-slip guard)."""
        target = os.path.realpath(os.path.join(base, *parts))
        base_real = os.path.realpath(base)
        if not target.startswith(base_real + os.sep) and target != base_real:
            raise ValueError(f"Path escapes extraction directory: {os.path.join(*parts)}")
        return target

    def _parse_xml_safe(path: str) -> ET.ElementTree:
        """Parse XML. Python's ElementTree never resolves external entities, so
        it is not vulnerable to XXE; this wrapper documents that guarantee."""
        return ET.parse(path)

    try:
        # Extract EPUB (it's a ZIP) — sanitize every entry path
        epub_path = _safe_join(epub_dir, "book.epub")
        with open(epub_path, "wb") as f:
            f.write(content)

        with zipfile.ZipFile(epub_path, "r") as zf:
            for member in zf.infolist():
                member_path = _safe_join(epub_dir, member.filename)
                if member.is_dir():
                    os.makedirs(member_path, exist_ok=True)
                    continue
                os.makedirs(os.path.dirname(member_path), exist_ok=True)
                with zf.open(member) as src, open(member_path, "wb") as dst:
                    shutil.copyfileobj(src, dst, length=1024 * 1024)

        # Parse container.xml to find the OPF
        container_path = os.path.join(epub_dir, "META-INF", "container.xml")
        if not os.path.exists(container_path):
            return RedirectResponse(url=f"/admin/books/{book_id}?error=invalid_epub", status_code=302)

        tree = _parse_xml_safe(container_path)
        rootfs = tree.getroot()
        ns = {"c": "urn:oasis:names:tc:opendocument:xmlns:container"}
        rootfile = rootfs.find(".//c:rootfile", ns)
        if rootfile is None:
            return RedirectResponse(url=f"/admin/books/{book_id}?error=no_opf", status_code=302)

        opf_path = _safe_join(epub_dir, rootfile.get("full-path", ""))
        if not os.path.exists(opf_path):
            return RedirectResponse(url=f"/admin/books/{book_id}?error=opf_not_found", status_code=302)

        result = await db.execute(select(Book).where(Book.id == book_id))
        book = result.scalar_one_or_none()
        if not book:
            return RedirectResponse(url="/admin/books", status_code=302)

        # Parse OPF
        opf_dir = os.path.dirname(opf_path)
        opf_tree = _parse_xml_safe(opf_path)
        opf_root = opf_tree.getroot()
        pkg_ns = {"p": "http://www.idpf.org/2007/opf"}

        spine = _find_el(opf_root, ".//p:spine", "spine", pkg_ns)
        if spine is None:
            return RedirectResponse(url=f"/admin/books/{book_id}?error=no_spine", status_code=302)

        manifest = {}
        for item in (opf_root.findall(".//p:item", pkg_ns) or opf_root.findall("item")):
            item_id = item.get("id", "")
            item_href = item.get("href", "")
            manifest[item_id] = item_href

        # Update book metadata from OPF
        metadata = _find_el(opf_root, ".//p:metadata", "metadata", pkg_ns)
        if metadata is not None:
            dc_ns = {"dc": "http://purl.org/dc/elements/1.1/"}
            title_el = _find_el(metadata, "dc:title", "title", dc_ns)
            if title_el is not None and title_el.text:
                book.title = title_el.text[:255]
            author_el = _find_el(metadata, "dc:creator", "creator", dc_ns)
            if author_el is not None and author_el.text:
                book.author = author_el.text[:255]
            publisher_el = _find_el(metadata, "dc:publisher", "publisher", dc_ns)
            if publisher_el is not None and publisher_el.text:
                book.publisher = publisher_el.text[:255]

        # Extract cover image if available
        cover_href = None
        for meta in (opf_root.findall(".//p:meta", pkg_ns) or opf_root.findall("meta")):
            if meta.get("name", "").lower() == "cover":
                cover_id = meta.get("content", "")
                if cover_id in manifest:
                    cover_href = manifest[cover_id]
                    break

        if not cover_href:
            for item_id, href in manifest.items():
                if "cover" in href.lower() and href.endswith((".jpg", ".png", ".jpeg")):
                    cover_href = href
                    break

        if cover_href:
            cover_path = _safe_join(opf_dir, cover_href)
            if os.path.exists(cover_path):
                ext = os.path.splitext(cover_href)[1]
                cover_filename = f"{book_id}{ext}"
                cover_dest = os.path.join(COVERS_DIR, cover_filename)
                shutil.copy2(cover_path, cover_dest)
                book.cover_url = f"/media/covers/{cover_filename}"

        # Delete existing chapters
        existing = await db.execute(select(Chapter).where(Chapter.book_id == book_id))
        for ch in existing.scalars().all():
            await db.delete(ch)

        # Parse each spine item as a chapter
        word_count = 0
        word_id_counter = 0
        for idx, itemref in enumerate(spine):
            idref = itemref.get("idref", "")
            if idref not in manifest:
                continue
            href = manifest[idref]
            if not href.endswith((".html", ".xhtml", ".htm", ".xml")):
                continue

            content_path = _safe_join(opf_dir, href)
            if not os.path.exists(content_path):
                continue

            with open(content_path, "r", encoding="utf-8", errors="replace") as fh:
                html_content = fh.read(5 * 1024 * 1024)  # 5 MB cap per chapter

            soup = _make_soup(html_content)
            paragraphs = []
            for tag in soup.find_all(["p", "h1", "h2", "h3", "h4", "h5", "h6", "div"]):
                text = tag.get_text(strip=True)
                if not text:
                    continue
                words = text.split()
                if len(words) < 3:
                    continue
                para_words = []
                for w in words:
                    wid = f"w{word_id_counter:06d}"
                    para_words.append({"id": wid, "text": w})
                    word_id_counter += 1
                paragraphs.append({"words": para_words})

            if not paragraphs:
                continue

            word_count += sum(len(p["words"]) for p in paragraphs)

            chapter_title = f"Chapter {idx + 1}"
            title_tag = soup.find("title")
            if title_tag and title_tag.get_text(strip=True):
                chapter_title = title_tag.get_text(strip=True)[:255]

            chapter = Chapter(
                id=str(uuid.uuid4()),
                book_id=book_id,
                title=chapter_title,
                order_index=idx,
                content=paragraphs,
            )
            db.add(chapter)

        if not book.title or book.title == "Unknown":
            book.title = "Unknown Title"
        if not book.author or book.author == "Unknown":
            book.author = "Unknown Author"

        book.word_count = word_count
        book.status = BookStatus.PUBLISHED
        await db.commit()

    except Exception as e:
        # Surface the failure to the admin instead of pretending it worked.
        logger.exception("EPUB import failed for book %s", book_id)
        await db.rollback()
        return RedirectResponse(
            url=f"/admin/books/{book_id}?error=epub_import_failed",
            status_code=302,
        )
    finally:
        shutil.rmtree(epub_dir, ignore_errors=True)

    return RedirectResponse(url=f"/admin/books/{book_id}?uploaded=epub", status_code=302)

# ===== USERS =====

@router.get("/users", response_class=HTMLResponse)
async def users_page(request: Request, search: str = "", page_no: int = 1,
                     db: AsyncSession = Depends(get_db)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    PAGE_SIZE = 25
    page = max(1, page_no or 1)
    base = select(User)
    if search and search.strip():
        escaped = search.strip().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        base = base.where(or_(
            User.email.ilike(f"%{escaped}%", escape="\\"),
            User.phone.ilike(f"%{escaped}%", escape="\\"),
        ))
    total = await db.scalar(select(func.count()).select_from(base.subquery())) or 0
    users = (await db.execute(
        base.order_by(User.created_at.desc())
        .offset((page - 1) * PAGE_SIZE)
        .limit(PAGE_SIZE)
    )).scalars().all()
    pages = max(1, -(-total // PAGE_SIZE))
    ctx = {
        "request": request,
        "page": "users",
        "users": [
            {
                "id": u.id,
                "email": u.email,
                "phone": u.phone,
                "is_active": u.is_active,
                "is_admin": u.is_admin,
                "is_verified": u.is_verified,
                "phone_verified": u.phone_verified,
                "created_at": u.created_at,
            }
            for u in users
        ],
        "total": total,
        "page_no": page,
        "pages": pages,
        "search": search or "",
    }
    return templates.TemplateResponse(request, "admin/users.html", await _inject_csrf(ctx, request))

@router.post("/users/create")
async def create_user(request: Request, email: str = Form(...), password: str = Form(...),
                      is_admin: bool = Form(False),
                      db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    existing = await db.execute(select(User).where(User.email == email))
    if existing.scalar_one_or_none():
        return RedirectResponse(url="/admin/users?error=exists", status_code=302)
    user = User(id=str(uuid.uuid4()), email=email, hashed_password=get_password_hash(password),
                is_verified=True, is_admin=is_admin)
    db.add(user)
    await db.commit()
    return RedirectResponse(url="/admin/users", status_code=302)

@router.post("/users/{user_id}/toggle")
async def toggle_user(user_id: str, request: Request,
                      db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(User).where(User.id == user_id))
    user = result.scalar_one_or_none()
    if user:
        user.is_active = not user.is_active
        await db.commit()
    return RedirectResponse(url="/admin/users", status_code=302)

@router.post("/users/{user_id}/delete")
async def delete_user_route(user_id: str, request: Request,
                            db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(User).where(User.id == user_id))
    user = result.scalar_one_or_none()
    if user and not user.is_admin:
        await db.delete(user)
        await db.commit()
    return RedirectResponse(url="/admin/users", status_code=302)

# ===== SUBSCRIPTIONS =====

@router.get("/subscriptions", response_class=HTMLResponse)
async def subscriptions_page(request: Request, db: AsyncSession = Depends(get_db)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    plans = (await db.execute(select(SubscriptionPlan).order_by(SubscriptionPlan.price))).scalars().all()
    # Join user + plan so the table shows emails and plan names, not raw UUIDs.
    sub_rows = (await db.execute(
        select(UserSubscription, User.email, SubscriptionPlan.name)
        .join(User, User.id == UserSubscription.user_id)
        .join(SubscriptionPlan, SubscriptionPlan.id == UserSubscription.plan_id, isouter=True)
        .order_by(UserSubscription.created_at.desc())
        .limit(50)
    )).all()
    user_subs = [
        {
            "email": email,
            "plan_name": plan_name or "—",
            "status": sub.status,
            "started_at": sub.started_at,
            "expires_at": sub.expires_at,
        }
        for sub, email, plan_name in sub_rows
    ]
    ctx = {"request": request, "page": "subscriptions", "plans": plans, "user_subs": user_subs}
    return templates.TemplateResponse(request, "admin/subscriptions.html", await _inject_csrf(ctx, request))

@router.post("/subscriptions/plan")
async def create_plan(request: Request, name: str = Form(...), description: str = Form(""),
                      price: float = Form(...), interval: str = Form("monthly"),
                      db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    plan = SubscriptionPlan(id=str(uuid.uuid4()), name=name, description=description,
                            price=price, interval=interval, is_active=True)
    db.add(plan)
    await db.commit()
    return RedirectResponse(url="/admin/subscriptions", status_code=302)

@router.post("/subscriptions/plan/{plan_id}/toggle")
async def toggle_plan(plan_id: str, request: Request,
                      db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(SubscriptionPlan).where(SubscriptionPlan.id == plan_id))
    plan = result.scalar_one_or_none()
    if plan:
        plan.is_active = not plan.is_active
        await db.commit()
    return RedirectResponse(url="/admin/subscriptions", status_code=302)

# ===== CATEGORIES =====

@router.get("/categories", response_class=HTMLResponse)
async def categories_page(request: Request, db: AsyncSession = Depends(get_db)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    cats = (await db.execute(select(Category).order_by(Category.name))).scalars().all()
    # Count books per category so the list shows real usage.
    counts = dict((await db.execute(
        select(BookCategory.category_id, func.count(BookCategory.book_id))
        .group_by(BookCategory.category_id)
    )).all())
    categories = [
        {
            "id": c.id,
            "name": c.name,
            "slug": c.slug,
            "description": c.description,
            "book_count": counts.get(c.id, 0),
        }
        for c in cats
    ]
    ctx = {"request": request, "page": "categories", "categories": categories}
    return templates.TemplateResponse(request, "admin/categories.html", await _inject_csrf(ctx, request))

@router.post("/categories")
async def create_category(request: Request, name: str = Form(...), description: str = Form(""),
                          db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    name = name.strip()
    slug = "-".join(name.lower().split())
    # Friendly duplicate handling — name and slug are both unique.
    existing = (await db.execute(
        select(Category).where(or_(Category.name == name, Category.slug == slug))
    )).scalar_one_or_none()
    if existing:
        return RedirectResponse(
            url="/admin/categories?error=duplicate", status_code=302
        )

    cat = Category(id=str(uuid.uuid4()), name=name, slug=slug,
                   description=description or None)
    db.add(cat)
    try:
        await db.commit()
    except IntegrityError:
        # Lost a concurrent race on the unique name/slug.
        await db.rollback()
        return RedirectResponse(url="/admin/categories?error=duplicate", status_code=302)
    return RedirectResponse(url="/admin/categories?created=1", status_code=302)

@router.post("/categories/{cat_id}/delete")
async def delete_category(cat_id: str, request: Request,
                          db: AsyncSession = Depends(get_db), _: bool = Depends(_verify_csrf)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    result = await db.execute(select(Category).where(Category.id == cat_id))
    cat = result.scalar_one_or_none()
    if cat:
        await db.delete(cat)
        await db.commit()
    return RedirectResponse(url="/admin/categories", status_code=302)

# ===== ANALYTICS =====

@router.get("/analytics", response_class=HTMLResponse)
async def analytics_page(request: Request, db: AsyncSession = Depends(get_db)):
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")
    days = 30
    since = datetime.now(timezone.utc) - timedelta(days=days)

    stats = {
        "users": (await db.scalar(select(func.count(User.id)))) or 0,
        "new_users": (await db.scalar(
            select(func.count(User.id)).where(User.created_at >= since)
        )) or 0,
        "books": (await db.scalar(select(func.count(Book.id)))) or 0,
        "subscriptions": (await db.scalar(
            select(func.count(UserSubscription.id)).where(UserSubscription.status == "active")
        )) or 0,
    }

    # Revenue
    stats["revenue"] = int(await db.scalar(
        select(func.coalesce(func.sum(Payment.amount), 0))
        .where(Payment.status == "completed")
    ) or 0)
    stats["revenue_month"] = int(await db.scalar(
        select(func.coalesce(func.sum(Payment.amount), 0))
        .where(Payment.status == "completed", Payment.completed_at >= since)
    ) or 0)
    stats["payments"] = await db.scalar(
        select(func.count(Payment.id)).where(Payment.status == "completed")
    ) or 0
    stats["payments_month"] = await db.scalar(
        select(func.count(Payment.id))
        .where(Payment.status == "completed", Payment.completed_at >= since)
    ) or 0
    stats["avg_order"] = int(stats["revenue"] / stats["payments"]) if stats["payments"] else 0

    # Revenue split by payment method
    stats["by_method"] = [
        {"method": m or "unknown", "total": float(t or 0), "count": int(c or 0)}
        for m, t, c in (await db.execute(
            select(Payment.method, func.sum(Payment.amount), func.count(Payment.id))
            .where(Payment.status == "completed", Payment.completed_at >= since)
            .group_by(Payment.method)
            .order_by(func.sum(Payment.amount).desc())
        )).all()
    ]

    # Reading engagement
    stats["reading_sessions"] = await db.scalar(
        select(func.count(ReadingSession.id)).where(ReadingSession.date >= since.date())
    ) or 0
    stats["reading_minutes"] = int((await db.scalar(
        select(func.coalesce(func.sum(ReadingSession.duration_seconds), 0))
        .where(ReadingSession.date >= since.date())
    ) or 0) / 60)

    # Most-read books (by session volume), falling back to newest when empty.
    popular = (await db.execute(
        select(Book.id, Book.title, Book.author, Book.cover_url, Book.updated_at,
               func.count(ReadingSession.id).label("sessions"))
        .join(ReadingSession, ReadingSession.book_id == Book.id)
        .group_by(Book.id, Book.title, Book.author, Book.cover_url, Book.updated_at)
        .order_by(func.count(ReadingSession.id).desc())
        .limit(8)
    )).all()
    if not popular:
        newest = (await db.execute(
            select(Book.id, Book.title, Book.author, Book.cover_url, Book.updated_at)
            .order_by(Book.created_at.desc())
            .limit(8)
        )).all()
        stats["popular"] = [
            {"id": r.id, "title": r.title, "author": r.author,
             "cover_url": r.cover_url, "updated_at": r.updated_at, "sessions": 0}
            for r in newest
        ]
    else:
        stats["popular"] = [
            {"id": r.id, "title": r.title, "author": r.author,
             "cover_url": r.cover_url, "updated_at": r.updated_at, "sessions": r.sessions}
            for r in popular
        ]

    ctx = {"request": request, "page": "analytics", "stats": stats, "days": days}
    return templates.TemplateResponse(request, "admin/analytics.html", await _inject_csrf(ctx, request))

@router.get("/analytics/export")
async def export_analytics(request: Request, db: AsyncSession = Depends(get_db)):
    """Export analytics to an Excel workbook (.xlsx) - FRS §13."""
    admin = await _get_admin(request, db)
    if not admin:
        return RedirectResponse(url="/admin/login")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    total_users = (await db.scalar(select(func.count(User.id)))) or 0
    total_books = (await db.scalar(select(func.count(Book.id)))) or 0
    active_subscriptions = (await db.scalar(
        select(func.count(UserSubscription.id)).where(UserSubscription.status == "active")
    )) or 0
    payments = (await db.execute(
        select(Payment).where(Payment.status == "completed").order_by(Payment.created_at)
    )).scalars().all()
    total_revenue = sum(p.amount for p in payments)
    books = (await db.execute(select(Book).order_by(Book.created_at.desc()).limit(10))).scalars().all()

    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["LYRR Analytics", ""])
    ws.append(["Generated", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")])
    ws.append([])
    ws.append(["Metric", "Value"])
    ws.append(["Total Users", total_users])
    ws.append(["Total Books", total_books])
    ws.append(["Active Subscriptions", active_subscriptions])
    ws.append(["Total Revenue", total_revenue])
    ws.append(["Completed Payments", len(payments)])

    ws2 = wb.create_sheet("Payments")
    ws2.append(["ID", "Amount", "Currency", "Method", "Status", "Date"])
    for p in payments:
        ws2.append([p.id, p.amount, p.currency, p.method, p.status, str(p.created_at)[:19]])

    ws3 = wb.create_sheet("Recent Books")
    ws3.append(["Title", "Author", "Status"])
    for b in books:
        ws3.append([b.title, b.author, b.status.value if hasattr(b.status, "value") else b.status])

    header_font = Font(bold=True)
    fill = PatternFill(start_color="DDEEFF", end_color="DDEEFF", fill_type="solid")
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = fill
        sheet.column_dimensions[get_column_letter(1)].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"lyrr_analytics_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
